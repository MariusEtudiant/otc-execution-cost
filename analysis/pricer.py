"""Write the Excel clip pricer (reports/otc_clip_pricer.xlsx) with live formulas.

Data sheets hold the measured sweep costs and the six-month liquidity clock; the Pricer,
Quote grid and Cost of waiting sheets are formulas only, so a trader can change the inputs
(asset, side, size, hour, margin, participation) and everything recomputes in Excel.
"""

import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS = os.path.join(ROOT, "reports")
SIZES = [0.1, 0.5, 1, 2, 5, 10]
WINDOWS = ["All measured", "London morning", "London + New York", "New York afternoon"]

BLUE = "2A78D6"
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
HEAD_FILL = PatternFill("solid", fgColor="EEF4FC")
OUT_FILL = PatternFill("solid", fgColor="E8F6EF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
MUTED = Font(color="6B6A66", italic=True, size=9)
THIN = Border(bottom=Side(style="thin", color="D9D8D4"))


def window_of_hour(h):
    if 7 <= h < 12:
        return "London morning"
    if 12 <= h < 16:
        return "London + New York"
    if 16 <= h < 21:
        return "New York afternoon"
    return "All measured"


def cost_table(m):
    """Median sweep cost (avg buy/sell, bps) by asset, venue, window and size; only windows with >= 60 minutes."""
    rows = []
    for (asset, venue), d in m.groupby(["asset", "venue"]):
        groups = [("All measured", d)] + [(w, g) for w, g in d.groupby("window") if g["ts_utc"].nunique() >= 60]
        for window, g in groups:
            row = {"key": f"{asset}|{venue}|{window}", "asset": asset, "venue": venue, "window": window,
                   "minutes": g["ts_utc"].nunique()}
            for s in SIZES:
                row[s] = round(g[f"cost_{s}m"].median(), 3)
            rows.append(row)
    return pd.DataFrame(rows)


def write_table(ws, df, start_row=1):
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(start_row, j, col if not isinstance(col, float) else col)
        c.font, c.fill = BOLD, HEAD_FILL
    for i, rec in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(rec, 1):
            ws.cell(i, j, None if pd.isna(v) else v)
    for j in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.column_dimensions["A"].width = 34
    ws.freeze_panes = ws.cell(start_row + 1, 2)


def build(m, clocks, last_mid):
    wb = Workbook()
    pricer = wb.active
    pricer.title = "Pricer"
    grid = wb.create_sheet("Quote grid")
    waiting = wb.create_sheet("Cost of waiting")
    sweep = wb.create_sheet("Sweep costs")
    clock = wb.create_sheet("Liquidity clock")
    notes = wb.create_sheet("Notes")

    # ---------------------------------------------------------------- data sheets
    costs = cost_table(m)
    write_table(sweep, costs)
    n_cost = len(costs) + 1
    sweep.cell(n_cost + 2, 1, "Median cost of sweeping the displayed book, bps vs the venue's own mid, "
                              "average of buy and sell. Columns are clip sizes in M$.").font = MUTED

    clock_rows = []
    for asset, c in clocks.items():
        for (weekend, hour), r in c.iterrows():
            day = "Weekend" if weekend else "Weekday"
            clock_rows.append({"key": f"{asset}|{day}|{hour}", "asset": asset, "day": day, "hour": hour,
                               "vol_per_min_musd": round(r["vol_per_min"] / 1e6, 4),
                               "sigma_1m_bps": round(r["sigma_1m"], 3),
                               "coinbase_share": round(r["coinbase_share"], 3), "days": int(r["days"])})
    clock_df = pd.DataFrame(clock_rows)
    write_table(clock, clock_df)
    n_clock = len(clock_df) + 1
    clock.cell(n_clock + 2, 1, "Median over days of the hourly mean of Binance + Coinbase volume (M$ per minute) "
                               "and of the RMS 1-minute log return (bps). 1 Apr to 12 Oct 2026.").font = MUTED

    sizes_rng = f"'Sweep costs'!$F$1:$K$1"
    keys_rng = f"'Sweep costs'!$A$2:$A${n_cost}"
    vals_rng = f"'Sweep costs'!$F$2:$K${n_cost}"
    ckeys = f"'Liquidity clock'!$A$2:$A${n_clock}"
    cvol = f"'Liquidity clock'!$E$2:$E${n_clock}"
    csig = f"'Liquidity clock'!$F$2:$F${n_clock}"

    # ---------------------------------------------------------------- pricer
    p = pricer
    p.column_dimensions["A"].width = 38
    p.column_dimensions["B"].width = 20
    p.column_dimensions["C"].width = 70
    p["A1"] = "OTC clip pricer: BTC and ETH"
    p["A1"].font = TITLE
    p["A2"] = "Yellow cells are inputs. Everything else is a formula on the data sheets."
    p["A2"].font = MUTED

    inputs = [
        ("Asset", "BTC", "BTC or ETH"),
        ("Client side", "Client buys", "Client buys: the desk sells and hedges by buying"),
        ("Clip size (M$)", 5, "0.1 to 10 M$ (sizes outside are clipped)"),
        ("Hour of the request (UTC)", 9, "0 to 23"),
        ("Day", "Weekday", "Weekday or Weekend"),
        ("Reference mid ($)", None, "Last mid recorded; overwrite with the live price"),
        ("Desk margin (bps)", 2, "Commercial margin added on top of the hedge cost"),
        ("Participation rate if worked", 0.10, "Share of market volume the desk would take to work the hedge"),
        ("Confidence (z)", 1.65, "1.65 = one-sided 95% on the inventory risk"),
    ]
    r0 = 4
    p.cell(r0, 1, "Inputs").font = BOLD
    for i, (label, value, note) in enumerate(inputs, r0 + 1):
        p.cell(i, 1, label)
        c = p.cell(i, 2, value)
        c.fill = INPUT_FILL
        p.cell(i, 3, note).font = MUTED
    asset, side, size, hour, day, mid, margin, part, z = (f"$B${r0 + 1 + k}" for k in range(9))
    p[mid[1:].replace("$", "")] = f'=IF({asset}="BTC",{last_mid["BTC"]},{last_mid["ETH"]})'
    p[mid[1:].replace("$", "")].fill = INPUT_FILL
    p[mid[1:].replace("$", "")].number_format = "#,##0.00"
    p[part[1:].replace("$", "")].number_format = "0%"

    dv = DataValidation(type="list", formula1='"BTC,ETH"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"Client buys,Client sells"', allow_blank=False)
    dv3 = DataValidation(type="list", formula1='"Weekday,Weekend"', allow_blank=False)
    dv4 = DataValidation(type="whole", operator="between", formula1="0", formula2="23")
    for d, cell in ((dv, asset), (dv2, side), (dv3, day), (dv4, hour)):
        p.add_data_validation(d)
        d.add(cell.replace("$", ""))

    r1 = r0 + len(inputs) + 2
    p.cell(r1, 1, "Calculation").font = BOLD
    n = f"MAX(0.1,MIN(10,{size}))"
    window_formula = (f'IF(AND({hour}>=7,{hour}<12),"London morning",IF(AND({hour}>=12,{hour}<16),'
                      f'"London + New York",IF(AND({hour}>=16,{hour}<21),"New York afternoon","All measured")))')
    calc = [
        ("Measured window for this hour", None, "Falls back to all measured minutes if this window was not collected"),
        ("Sweep curve used", None, "Combined book of Binance, Coinbase and Kraken"),
        ("Size bracket (index)", None, "Position of the clip between the measured sizes"),
        ("Sweep cost, combined books (bps)", None, "Linear interpolation between measured sizes"),
        ("Sweep cost, best single venue (bps)", None, "Cheapest of the three venues on its own"),
        ("Volume at this hour (M$ per minute)", None, "Binance + Coinbase, six-month median"),
        ("Volatility at this hour (bps per minute)", None, "Six-month median"),
        ("Minutes to work the clip", None, "Clip / (participation x volume)"),
        ("Risk of working it (bps)", None, "z x sigma x sqrt(T / 3): linear liquidation, Almgren-Chriss"),
        ("Hedge cost retained (bps)", None, "Cheaper of sweeping now or the risk of working it"),
        ("Hedge method", None, ""),
    ]
    for i, (label, _, note) in enumerate(calc, r1 + 1):
        p.cell(i, 1, label)
        p.cell(i, 3, note).font = MUTED
    rw, rkey, ridx, rsweep, rsingle, rvol, rsig, rmin, rrisk, rhedge, rmethod = (r1 + 1 + k for k in range(11))
    key_all = f'{asset}&"|ALL3|"&B{rw}'
    p[f"B{rw}"] = (f'=IF(ISNUMBER(MATCH({asset}&"|ALL3|"&{window_formula},{keys_rng},0)),'
                   f'{window_formula},"All measured")')
    p[f"B{rkey}"] = f"={key_all}"
    p[f"B{ridx}"] = f"=MIN(MATCH({n},{sizes_rng},1),5)"

    def interp(key_expr):
        row = f"MATCH({key_expr},{keys_rng},0)"
        x0 = f"INDEX({sizes_rng},1,B{ridx})"
        x1 = f"INDEX({sizes_rng},1,B{ridx}+1)"
        y0 = f"INDEX({vals_rng},{row},B{ridx})"
        y1 = f"INDEX({vals_rng},{row},B{ridx}+1)"
        return f"{y0}+({n}-{x0})*({y1}-{y0})/({x1}-{x0})"

    p[f"B{rsweep}"] = f"={interp(f'B{rkey}')}"
    singles = ",".join(interp(f'{asset}&"|{v}|"&B{rw}') for v in ("binance", "coinbase", "kraken"))
    p[f"B{rsingle}"] = f"=MIN({singles})"
    ckey = f'{asset}&"|"&{day}&"|"&{hour}'
    p[f"B{rvol}"] = f"=INDEX({cvol},MATCH({ckey},{ckeys},0))"
    p[f"B{rsig}"] = f"=INDEX({csig},MATCH({ckey},{ckeys},0))"
    p[f"B{rmin}"] = f"={n}/({part}*B{rvol})"
    p[f"B{rrisk}"] = f"={z}*B{rsig}*SQRT(B{rmin}/3)"
    p[f"B{rhedge}"] = f"=MIN(B{rsweep},B{rrisk})"
    p[f"B{rmethod}"] = f'=IF(B{rsweep}<=B{rrisk},"Sweep the combined books now","Work the order")'
    for r in (rsweep, rsingle, rvol, rsig, rmin, rrisk, rhedge):
        p[f"B{r}"].number_format = "0.00"

    r2 = rmethod + 2
    p.cell(r2, 1, "Quote").font = BOLD
    out = [
        ("Quoted spread to the client (bps)", f"=B{rhedge}+{margin}", "0.00"),
        ("Client price ($)", f'=IF({side}="Client buys",{mid}*(1+B{r2 + 1}/10000),{mid}*(1-B{r2 + 1}/10000))',
         "#,##0.00"),
        ("Expected desk P&L on the clip ($)", f"={margin}/10000*{n}*1000000", "#,##0"),
        ("Saving vs hedging on the best single venue ($)", f"=(B{rsingle}-B{rsweep})/10000*{n}*1000000", "#,##0"),
    ]
    for i, (label, formula, fmt) in enumerate(out, r2 + 1):
        p.cell(i, 1, label).font = BOLD
        c = p.cell(i, 2, formula)
        c.number_format, c.fill, c.font = fmt, OUT_FILL, BOLD
    p.cell(r2 + 6, 1, "Sweep costs come from minute snapshots of the displayed books (no hidden liquidity, no refill, "
                      "no fees). Read them as the cost of an aggressive hedge, not a guaranteed fill.").font = MUTED

    # ---------------------------------------------------------------- quote grid
    g = grid
    g["A1"] = "Quote grid: spread to quote for an immediate hedge on the combined books (bps)"
    g["A1"].font = TITLE
    g["A2"] = "Spread = measured sweep cost + desk margin (Pricer!B" + margin[3:] + "). Asset follows the Pricer."
    g["A2"].font = MUTED
    g["A4"], g["A4"].font = "Clip (M$)", BOLD
    for j, w in enumerate(WINDOWS, 2):
        c = g.cell(4, j, w)
        c.font, c.fill, c.alignment = BOLD, HEAD_FILL, Alignment(wrap_text=True, horizontal="center")
        g.column_dimensions[get_column_letter(j)].width = 18
    g.column_dimensions["A"].width = 14
    for i, s in enumerate(SIZES, 5):
        g.cell(i, 1, s).font = BOLD
        for j in range(2, 2 + len(WINDOWS)):
            col = get_column_letter(j)
            key = f'Pricer!{asset}&"|ALL3|"&{col}$4'
            c = g.cell(i, j, f'=IFERROR(INDEX({vals_rng},MATCH({key},{keys_rng},0),'
                             f'MATCH($A{i},{sizes_rng},0))+Pricer!{margin},"not measured")')
            c.number_format = "0.0"
            c.alignment = Alignment(horizontal="center")
    last = 4 + len(SIZES)
    g.conditional_formatting.add(f"B5:{get_column_letter(1 + len(WINDOWS))}{last}",
                                 ColorScaleRule(start_type="min", start_color="F3F8FE",
                                                end_type="max", end_color="5598E7"))

    # ---------------------------------------------------------------- cost of waiting
    w = waiting
    w["A1"] = "Cost of waiting: 95% risk of working the clip instead of sweeping (bps)"
    w["A1"].font = TITLE
    w["A2"] = ("Clip size, participation and z follow the Pricer. Compare with the sweep cost in the Pricer "
               "(a few bps): waiting is the expensive option for every hour.")
    w["A2"].font = MUTED
    heads = ["Hour (UTC)", "BTC weekday", "BTC weekend", "ETH weekday", "ETH weekend"]
    for j, h in enumerate(heads, 1):
        c = w.cell(4, j, h)
        c.font, c.fill = BOLD, HEAD_FILL
        w.column_dimensions[get_column_letter(j)].width = 15
    for hr in range(24):
        i = 5 + hr
        w.cell(i, 1, hr).font = BOLD
        for j, (a, d) in enumerate((("BTC", "Weekday"), ("BTC", "Weekend"), ("ETH", "Weekday"), ("ETH", "Weekend")), 2):
            k = f'"{a}|{d}|"&$A{i}'
            vol = f"INDEX({cvol},MATCH({k},{ckeys},0))"
            sig = f"INDEX({csig},MATCH({k},{ckeys},0))"
            c = w.cell(i, j, f"=Pricer!{z}*{sig}*SQRT(MAX(0.1,MIN(10,Pricer!{size}))/(Pricer!{part}*{vol})/3)")
            c.number_format = "0.0"
    w.conditional_formatting.add("B5:E28", ColorScaleRule(start_type="min", start_color="F3F8FE",
                                                          end_type="max", end_color="5598E7"))

    # ---------------------------------------------------------------- notes
    lines = [
        ("How to read this workbook", True),
        ("Pricer: one client request in, one quote out. Change the yellow cells.", False),
        ("Quote grid: the same quote for every measured size and session.", False),
        ("Cost of waiting: what the desk would risk by working the clip at a share of volume instead of sweeping.", False),
        ("", False),
        ("Data", True),
        (f"Sweep costs: {m['ts_utc'].nunique()} one-minute snapshots of the Binance (BTCUSDT, ETHUSDT), Coinbase and "
         f"Kraken (BTC-USD, ETH-USD) books, {m['ts_utc'].min():%d %b %H:%M} to {m['ts_utc'].max():%d %b %H:%M} UTC.", False),
        ("Liquidity clock: Binance and Coinbase 1-minute candles, 1 April to 12 October 2026.", False),
        ("Costs are in bps of each venue's own mid, so the USDT/USD basis between Binance and the USD venues cancels out.", False),
        ("", False),
        ("Limits", True),
        ("Displayed liquidity only: no hidden orders, no book refill after a sweep, no fees, no market-maker fade.", False),
        ("Asia was not measured in the books; for those hours the Pricer uses all measured minutes.", False),
        ("Volume excludes Kraken and every other venue, so the time to work a clip is conservative.", False),
    ]
    notes.column_dimensions["A"].width = 120
    for i, (text, bold) in enumerate(lines, 1):
        notes.cell(i, 1, text).font = BOLD if bold else Font()

    for ws in (pricer, grid, waiting):
        ws.sheet_view.showGridLines = False
    path = os.path.join(REPORTS, "otc_clip_pricer.xlsx")
    wb.save(path)
    return path


if __name__ == "__main__":
    import sys

    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from build import load_history, liquidity_clock, load_snapshots

    m = load_snapshots()
    last = m[m.venue == "coinbase"].sort_values("ts_utc").groupby("asset")["mid"].last().round(2).to_dict()
    clocks = {a: liquidity_clock(load_history(a)) for a in ("BTC", "ETH")}
    print(build(m, clocks, last))
